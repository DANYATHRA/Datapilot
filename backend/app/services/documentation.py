"""Documentation & presentation service.

- Auto-generated data dictionary
- Process documentation (pipeline steps, cleaning decisions, quality results)
- Insights deck (PowerPoint via python-pptx)
- Stakeholder summary
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import pandas as pd


# ── Data dictionary ────────────────────────────────────────────────

def generate_data_dictionary(df: pd.DataFrame, profile_data: dict,
                             dataset_name: str, semantics: dict | None = None) -> list[dict]:
    """Generate a per-column data dictionary.

    Each entry: column name, detected meaning, type, description (editable),
    allowed values, null %, example values (PII hidden), source, owner.
    """
    columns = profile_data.get("columns", [])
    pii_cols = set(profile_data.get("pii_columns", {}).keys())
    dictionary = []

    for col_info in columns:
        name = col_info["name"]
        is_pii = name in pii_cols

        # Example values (PII hidden)
        if is_pii:
            examples = ["[REDACTED]"] * 3
        else:
            vals = df[name].dropna().head(5).astype(str).tolist()
            examples = vals[:5]

        # Allowed values (for low-cardinality columns)
        allowed_values = None
        if col_info.get("kind") == "text" and col_info.get("unique", 0) <= 25:
            allowed_values = df[name].dropna().unique().tolist()[:25]

        sem_type = col_info.get("semantic_type", "")
        entry = {
            "column": name,
            "detected_meaning": sem_type,
            "data_type": col_info.get("dtype", ""),
            "kind": col_info.get("kind", ""),
            "description": _default_description(name, sem_type),
            "allowed_values": allowed_values,
            "null_pct": col_info.get("null_pct", 0),
            "null_count": col_info.get("nulls", 0),
            "unique_count": col_info.get("unique", 0),
            "example_values": examples,
            "is_pii": is_pii,
            "source": dataset_name,
            "owner": "",
        }
        dictionary.append(entry)

    return dictionary


def _default_description(name: str, sem_type: str) -> str:
    """Generate a default description from column name and semantic type."""
    clean = name.replace("_", " ").replace("-", " ").title()
    if sem_type and sem_type != "other_text":
        return f"{clean} ({sem_type.replace('_', ' ')})"
    return clean


def data_dictionary_html(dictionary: list[dict], dataset_name: str) -> str:
    """Export data dictionary as HTML."""
    html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>Data Dictionary – {dataset_name}</title>
<style>
  body {{ font-family: 'Segoe UI', system-ui, sans-serif; max-width: 1100px; margin: 2rem auto; padding: 0 1rem; color: #1a1a2e; }}
  h1 {{ color: #16213e; border-bottom: 2px solid #0f3460; padding-bottom: .5rem; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; font-size: .85rem; }}
  th, td {{ text-align: left; padding: .5rem .75rem; border-bottom: 1px solid #e2e8f0; }}
  th {{ background: #f8fafc; font-weight: 600; position: sticky; top: 0; }}
  .pii {{ color: #ef4444; font-weight: 600; }}
  .tag {{ display: inline-block; padding: .1rem .4rem; border-radius: .25rem; background: #e2e8f0; font-size: .75rem; margin: .1rem; }}
</style></head><body>
<h1>📖 Data Dictionary</h1>
<p><strong>Dataset:</strong> {dataset_name} &middot; <strong>Generated:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M')}</p>
<table>
<tr><th>Column</th><th>Type</th><th>Meaning</th><th>Description</th><th>Nulls</th><th>Unique</th><th>Examples</th></tr>"""

    for entry in dictionary:
        pii_badge = ' <span class="pii">[PII]</span>' if entry["is_pii"] else ""
        examples = ", ".join(str(e) for e in entry["example_values"][:3])
        html += f"""<tr>
<td><strong>{entry['column']}</strong>{pii_badge}</td>
<td>{entry['data_type']}</td>
<td><span class="tag">{entry['detected_meaning']}</span></td>
<td>{entry['description']}</td>
<td>{entry['null_pct']}%</td>
<td>{entry['unique_count']}</td>
<td style="max-width:200px;overflow:hidden;text-overflow:ellipsis">{examples}</td></tr>"""

    html += "</table></body></html>"
    return html


def data_dictionary_excel(dictionary: list[dict]) -> bytes:
    """Export data dictionary as an Excel file."""
    rows = []
    for entry in dictionary:
        rows.append({
            "Column": entry["column"],
            "Data Type": entry["data_type"],
            "Detected Meaning": entry["detected_meaning"],
            "Description": entry["description"],
            "Null %": entry["null_pct"],
            "Null Count": entry["null_count"],
            "Unique Values": entry["unique_count"],
            "Examples": ", ".join(str(e) for e in entry["example_values"][:5]),
            "Is PII": entry["is_pii"],
            "Source": entry["source"],
        })
    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data Dictionary")
    return buf.getvalue()


# ── Process documentation ──────────────────────────────────────────

def generate_process_doc(dataset_doc: dict, quality_report: dict | None = None) -> str:
    """Generate process documentation from the audit trail."""
    versions = dataset_doc.get("versions", [])
    name = dataset_doc.get("name", "Dataset")

    sections = [
        f"# Process Documentation — {name}\n",
        f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M')}\n",
        "## Data Pipeline Steps\n",
    ]

    for v in versions:
        step = v.get("step", {})
        op = step.get("op", "unknown")
        params = step.get("params", {})
        sections.append(f"### Version {v['v']}: {op.replace('_', ' ').title()}")
        sections.append(f"- **Rows:** {v.get('rows', 'N/A')}")
        sections.append(f"- **Columns:** {v.get('cols', 'N/A')}")
        if params:
            sections.append(f"- **Parameters:** {params}")
        sections.append("")

    if quality_report:
        sections.append("## Data Quality Results\n")
        sections.append(f"- **Overall Score:** {quality_report.get('overall_score', 'N/A')}%")
        dims = quality_report.get("dimensions", {})
        for dim, info in dims.items():
            sections.append(f"- **{dim.title()}:** {info.get('score', 'N/A')}%")
        sections.append("")

        issues = quality_report.get("issue_log", [])
        if issues:
            sections.append("## Known Issues\n")
            for issue in issues:
                sections.append(f"- [{issue['severity'].upper()}] {issue['column']}: {issue['rule_name']} "
                              f"({issue['fail_count']} failures)")

    sections.append("\n## Assumptions and Limitations\n")
    sections.append("- Data quality checks are based on automated rules; manual review recommended")
    sections.append("- Statistical outliers detected via IQR method; domain experts should validate")

    return "\n".join(sections)


# ── Stakeholder summary ───────────────────────────────────────────

def generate_stakeholder_summary(dataset_name: str, profile: dict,
                                 kpi_values: list[dict] | None = None,
                                 insights: list[dict] | None = None) -> str:
    """Generate a one-page executive summary for non-technical readers."""
    n_rows = profile.get("rows", 0)
    n_cols = profile.get("cols", 0)
    quality = profile.get("quality", {})
    score = quality.get("score", 0)

    summary = [
        f"# Executive Summary — {dataset_name}\n",
        f"**Date:** {datetime.now().strftime('%Y-%m-%d')}\n",
        "## Data Overview\n",
        f"The dataset contains **{n_rows:,} records** across **{n_cols} fields** "
        f"with an overall data quality score of **{score}%**.\n",
    ]

    if kpi_values:
        summary.append("## Key Performance Indicators\n")
        summary.append("| KPI | Value |")
        summary.append("|-----|-------|")
        for kpi in kpi_values:
            summary.append(f"| {kpi['name']} | {kpi['formatted']} |")
        summary.append("")

    if insights:
        summary.append("## Key Findings\n")
        for i, insight in enumerate(insights[:5], 1):
            text = insight.get("text", insight.get("description", ""))
            summary.append(f"{i}. {text}")
        summary.append("")

    summary.append("## Data Quality\n")
    summary.append(f"- Completeness: {quality.get('completeness', 'N/A')}%")
    summary.append(f"- Uniqueness: {quality.get('uniqueness', 'N/A')}%")
    summary.append(f"- Missing cells: {profile.get('missing_cells', 0):,}")
    summary.append(f"- Duplicate rows: {profile.get('duplicate_rows', 0):,}")

    summary.append("\n## Next Steps\n")
    summary.append("- Review data quality issues flagged in the detailed report")
    summary.append("- Validate KPI definitions with domain stakeholders")
    summary.append("- Schedule follow-up for deeper analysis of key trends")

    return "\n".join(summary)


# ── PowerPoint deck ────────────────────────────────────────────────

def generate_insights_deck(dataset_name: str, profile: dict,
                           kpi_values: list[dict] | None = None,
                           insights: list[dict] | None = None,
                           quality_report: dict | None = None,
                           chart_images: list[tuple[str, bytes]] | None = None) -> bytes:
    """Generate a PowerPoint insights deck.

    Returns: bytes of the .pptx file.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise ImportError("python-pptx is required for deck generation. Install with: pip install python-pptx")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color scheme
    DARK_BG = RGBColor(0x16, 0x21, 0x3E)
    PRIMARY = RGBColor(0x0F, 0x34, 0x60)
    ACCENT = RGBColor(0x53, 0x3C, 0xA6)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_GRAY = RGBColor(0xF0, 0xF0, 0xF5)

    def _add_title_slide(title: str, subtitle: str):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = DARK_BG
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(44)
        p.font.color.rgb = WHITE
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        # Subtitle
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = LIGHT_GRAY
        p2.alignment = PP_ALIGN.CENTER

    def _add_content_slide(title: str, bullets: list[str]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Title bar
        shape = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PRIMARY
        shape.line.fill.background()
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11), Inches(0.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.color.rgb = WHITE
        p.font.bold = True
        # Bullets
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(11), Inches(5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.text = bullet
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            p.space_after = Pt(12)

    # Slide 1: Title
    _add_title_slide(
        f"DataPilot Insights",
        f"{dataset_name} · {datetime.now().strftime('%B %Y')}"
    )

    # Slide 2: Objective & Data Summary
    quality = profile.get("quality", {})
    _add_content_slide("Objective & Data Summary", [
        f"Dataset: {dataset_name}",
        f"Records: {profile.get('rows', 0):,} rows × {profile.get('cols', 0)} columns",
        f"Data Quality Score: {quality.get('score', 0)}%",
        f"Completeness: {quality.get('completeness', 'N/A')}% | Uniqueness: {quality.get('uniqueness', 'N/A')}%",
        f"Missing cells: {profile.get('missing_cells', 0):,} | Duplicates: {profile.get('duplicate_rows', 0):,}",
    ])

    # Slide 3: KPIs
    if kpi_values:
        kpi_bullets = [f"{kpi['name']}: {kpi['formatted']}   ({kpi.get('definition', '')})" for kpi in kpi_values[:8]]
        _add_content_slide("Key Performance Indicators", kpi_bullets)

    # Slide 4-6: Insights (with evidence)
    if insights:
        for i in range(0, min(len(insights), 6), 2):
            batch = insights[i:i+2]
            bullets = []
            for ins in batch:
                text = ins.get("text", ins.get("description", ""))
                evidence = ins.get("evidence", ins.get("metric", ""))
                bullets.append(f"• {text}")
                if evidence:
                    bullets.append(f"  Evidence: {evidence}")
            _add_content_slide(f"Key Findings ({i//2 + 1})", bullets)

    # Slide: Quality Caveats
    caveats = [
        "Data quality checks are automated; manual validation recommended",
        f"Overall quality score: {quality.get('score', 0)}%",
    ]
    if quality_report and quality_report.get("issue_log"):
        for issue in quality_report["issue_log"][:3]:
            caveats.append(f"⚠ {issue['column']}: {issue['rule_name']} ({issue['fail_count']} failures)")
    _add_content_slide("Quality Caveats & Limitations", caveats)

    # Slide: Next Steps
    _add_content_slide("Next Steps", [
        "Review and validate KPI definitions with stakeholders",
        "Address critical data quality issues before production use",
        "Schedule follow-up analysis for identified trends",
        "Document data lineage and transformation decisions",
    ])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


generate_powerpoint_deck = generate_insights_deck
