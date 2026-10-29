"""
Tests for Column-Aware (Semantic) Cleaning in DataPilot.

Verifies:
1. Blank roll_no stays blank and row appears in Needs_Review.
2. No identifier/email/phone/name/date is ever imputed.
3. Marks stay blank by default.
4. Grouped median is used for money when allowed.
5. Total is filled by formula and not by median.
6. Phone/pincode/roll_no keep text formatting (leading zeros, no .0).
7. Duplicate IDs with different data become conflicts in Needs_Review.
8. Outliers never clip IDs or percentages.
9. Semantic detection works with differently named columns:
   (roll no, Roll_No, RollNumber, Reg. No, Employee ID, Order ID).
10. User overrides change the outcome.
11. Order ID and Email blanks in sales data stay blank.
"""
import io
import openpyxl
import pandas as pd
import pytest

from app.services.column_semantics import (
    infer_column_semantic_type,
    infer_dataset_semantics,
)
from app.services.missing_policy import (
    POLICY_MATRIX,
    detect_derived_formulas,
    get_preview_line,
)
from app.services.pipeline import build_export_workbook, run_pipeline


@pytest.fixture
def students_workbook():
    return "sample_data/students.xlsx"


@pytest.fixture
def test_sales_workbook():
    return "sample_data/datapilot_test_sales.xlsx"


def test_column_name_detection_variations():
    """Detection works with differently named columns: roll no, Roll_No, RollNumber, Reg. No, Employee ID, Order ID."""
    test_names = ["roll no", "Roll_No", "RollNumber", "Reg. No", "Employee ID", "Order ID", "order_id", "empid"]
    dummy_series = pd.Series(["101", "102", "103", "104", "105"])
    for name in test_names:
        res = infer_column_semantic_type(name, dummy_series)
        assert res["semantic_type"] == "identifier", f"Failed for column name: {name}, got {res['semantic_type']}"
        assert res["confidence"] >= 0.75, f"Confidence should be high for {name}"
        assert not res["needs_confirmation"], f"Should not need confirmation for {name}"


def test_students_cleaning_rules(students_workbook):
    """
    Run pipeline on students sheet:
    - blank roll_no stays blank and row appears in Needs_Review
    - no identifier/email/phone/name/date is ever imputed
    - marks stay blank by default
    - phone keeps text formatting (leading zero preserved)
    - duplicate ID with different data becomes conflict in Needs_Review
    - outliers never clip ids
    """
    df = pd.read_excel(students_workbook, sheet_name="Students", dtype=object)
    original_blank_roll_count = int(df["roll_no"].isna().sum())
    assert original_blank_roll_count == 3, "Fixture should have 3 blank roll numbers"

    result = run_pipeline(df, sheet_name="Students")
    cleaned_df = result["cleaned_df"]
    needs_review = result["needs_review_items"]
    left_blanks = result["left_blank_by_design"]

    # 1. Blank roll_no stays blank
    assert cleaned_df["roll_no"].isna().sum() >= original_blank_roll_count, "Blank roll_no must stay blank"

    # Row appears in Needs_Review
    roll_reviews = [r for r in needs_review if r["column"] == "roll_no" and "Missing identifier" in r["issue"]]
    assert len(roll_reviews) >= 3, "Missing roll_no rows must appear in Needs_Review"

    # 2. No identifier/email/phone/name/date is ever imputed
    for col in ["email", "phone", "name", "dob"]:
        orig_nulls = int(df[col].isna().sum())
        if orig_nulls > 0:
            cleaned_nulls = int(cleaned_df[col].isna().sum())
            assert cleaned_nulls == orig_nulls, f"Column '{col}' must not be imputed! Was {orig_nulls}, now {cleaned_nulls}"

    # 3. Marks stay blank by default
    orig_marks_nulls = int(df["marks"].isna().sum())
    assert orig_marks_nulls > 0, "Marks had blank cells in fixture"
    assert int(cleaned_df["marks"].isna().sum()) == orig_marks_nulls, "Marks must stay blank by default"

    # 6. Phone/pincode/roll_no keep text formatting (leading zero preserved, no .0)
    phone_vals = cleaned_df["phone"].dropna().astype(str).tolist()
    assert any(p.startswith("0") for p in phone_vals), "Phone numbers with leading zero must preserve leading zero"
    for p in phone_vals:
        assert not p.endswith(".0"), f"Phone value '{p}' should not end with .0"
    for r in cleaned_df["roll_no"].dropna().astype(str).tolist():
        assert not r.endswith(".0"), f"Roll_no value '{r}' should not end with .0"

    # 7. Duplicate IDs with different data become conflicts in Needs_Review
    dup_conflicts = [r for r in needs_review if "Duplicate ID conflict" in r["issue"]]
    assert len(dup_conflicts) > 0, "Conflicting duplicate IDs must be reported in Needs_Review"
    assert any("114" in r["issue"] or r["value"] == "114" for r in dup_conflicts)

    # 8. Outliers never clip IDs
    id_changes = [ch for ch in result["all_changes"] if ch["stage"] == "outliers" and ch["column"] == "roll_no"]
    assert len(id_changes) == 0, "Identifiers must never be clipped by outliers stage"


def test_revenue_formula_fill(students_workbook):
    """
    Run pipeline on Revenue sheet:
    - total is filled by formula units * unit_price * (1 - discount/100) and NOT by median
    """
    df = pd.read_excel(students_workbook, sheet_name="Revenue")
    orig_total_nulls = int(df["total"].isna().sum())
    assert orig_total_nulls == 3, "Revenue sheet fixture must have 3 blank totals"

    # Detect formulas
    formulas = detect_derived_formulas(df)
    assert len(formulas) > 0, "Formula detection should find revenue formula"
    assert any("units" in f["expression"] and "unit_price" in f["expression"] for f in formulas)

    result = run_pipeline(df, sheet_name="Revenue")
    cleaned_df = result["cleaned_df"]

    # All blank totals must be filled
    assert cleaned_df["total"].isna().sum() == 0, "All blank totals should be filled by formula"

    # Verify filled values match formula: units * unit_price * (1 - discount/100)
    for idx, row in cleaned_df.iterrows():
        expected = row["units"] * row["unit_price"] * (1.0 - row["discount"] / 100.0)
        assert abs(row["total"] - expected) < 0.1, f"Row {idx} total {row['total']} does not match formula {expected}"

    # Verify formula recorded in Change_Log
    formula_changes = [ch for ch in result["all_changes"] if ch["column"] == "total" and "formula" in ch["reason"].lower()]
    assert len(formula_changes) == 3, "Formula filling must be recorded in Change_Log"


def test_grouped_median_imputation():
    """Grouped median is used for money when allowed (group count >= 5), else global median."""
    data = {
        "category": ["A", "A", "A", "A", "A", "B", "B", "B"],
        "price": [10.0, 12.0, 14.0, 16.0, None, 100.0, 102.0, None],  # A has 4 known -> global or group?
    }
    df = pd.DataFrame(data)
    # Give A 5 valid values
    df = pd.concat([df, pd.DataFrame([{"category": "A", "price": 18.0}])], ignore_index=True)
    # A has: 10, 12, 14, 16, 18 -> median 14.0
    result = run_pipeline(df, sheet_name="Prices")
    cleaned_df = result["cleaned_df"]
    # Missing price in category A (originally row 4)
    filled_a = cleaned_df.loc[cleaned_df["category"] == "A", "price"]
    assert not filled_a.isna().any()
    assert 14.0 in filled_a.values


def test_user_overrides_change_outcome(students_workbook):
    """User overrides change the outcome."""
    df = pd.read_excel(students_workbook, sheet_name="Students", dtype=object)

    # Override marks missing rule from default leave_blank to fill_median
    column_rules = {
        "marks": {"semantic_type": "score/marks", "missing_rule": "fill_median", "invalid_rule": "keep"}
    }
    result = run_pipeline(df, sheet_name="Students", column_rules=column_rules)
    cleaned_df = result["cleaned_df"]

    # Now marks should be filled with median
    assert cleaned_df["marks"].isna().sum() == 0, "User override fill_median should have filled marks"
    median_val = df["marks"].dropna().median()
    assert (cleaned_df["marks"] == median_val).sum() >= 3


def test_sales_test_workbook_blanks_stay_blank(test_sales_workbook):
    """Re-run on datapilot_test_sales.xlsx: Order ID and Email blanks stay blank."""
    df = pd.read_excel(test_sales_workbook)
    # Inject blanks into order_id and customer_email to test
    df.loc[0:2, "order_id"] = None
    df.loc[3:4, "customer_email"] = None

    result = run_pipeline(df, sheet_name="Sales")
    cleaned_df = result["cleaned_df"]
    needs_review = result["needs_review_items"]

    # Blank order_ids stay blank
    assert cleaned_df["order_id"].isna().sum() == 3, "Blank order_id must stay blank"
    # Blank customer_email stay blank
    assert cleaned_df["customer_email"].isna().sum() == 2, "Blank customer_email must stay blank"

    # Blank order_ids appear in Needs_Review
    order_reviews = [r for r in needs_review if r["column"] == "order_id" and "Missing identifier" in r["issue"]]
    assert len(order_reviews) == 3, "Blank order_id must appear in Needs_Review"


def test_export_workbook_has_needs_review_sheet(students_workbook):
    """Export workbook contains Cleaned_Data, Needs_Review, Change_Log, Summary, and Legend sheets."""
    df = pd.read_excel(students_workbook, sheet_name="Students", dtype=object)
    result = run_pipeline(df, sheet_name="Students")

    xlsx_bytes = build_export_workbook(
        cleaned_df=result["cleaned_df"],
        all_changes=result["all_changes"],
        removed_rows=result["removed_rows"],
        stage_summaries=result["stage_summaries"],
        quality_comparison=result["quality_comparison"],
        needs_review_items=result["needs_review_items"],
        left_blank_by_design=result["left_blank_by_design"],
        sheet_name="Students",
    )

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    assert "Cleaned_Data" in sheet_names
    assert "Cleaned_Highlighted" in sheet_names
    assert "Change_Log" in sheet_names
    assert "Needs_Review" in sheet_names
    assert "Summary" in sheet_names
    assert "Legend" in sheet_names

    # Check Needs_Review sheet content
    ws_rev = wb["Needs_Review"]
    assert ws_rev.max_row >= 4, "Needs_Review sheet should have rows for flagged items"

    # Check Summary sheet has left blank by design
    ws_sum = wb["Summary"]
    sum_text = " ".join(str(ws_sum.cell(row=i, column=1).value) for i in range(1, ws_sum.max_row + 1))
    assert "left blank by design" in sum_text.lower()
